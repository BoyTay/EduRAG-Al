import io
import os
import tempfile
import unittest
from pathlib import Path

import pymupdf as fitz
from PIL import Image, ImageDraw
from docx import Document as DocxDocument

from document_ingestion import (
    annotate_chunk_metadata,
    detect_section_title,
    DocumentIngestionError,
    IngestionSettings,
    OCRResultCache,
    load_and_chunk_document,
)
from langchain_core.documents import Document
from ocr_service import OCRConfigurationError, OCRLine, OCRPageResult


class FakeOCR:
    def __init__(self, fingerprint="small-a", fail=False):
        self.calls = 0
        self.fail = fail
        self._fingerprint = fingerprint

    @property
    def cache_fingerprint(self):
        return {"provider": "fake", "version": self._fingerprint}

    def recognize_page(self, _image):
        self.calls += 1
        if self.fail:
            raise RuntimeError("fault injection")
        line = OCRLine("Điều 12. Nội dung nhận dạng từ trang scan.", 0.93, [5, 5, 100, 30])
        return OCRPageResult(line.text, line.confidence, [line], "fake-small")


class DocumentIngestionTests(unittest.TestCase):
    def test_invalid_ocr_alphabet_rejects_scan_with_page_number(self):
        class InvalidOCR(FakeOCR):
            def recognize_page(self, _image):
                raise OCRConfigurationError("Model thiếu ký tự tiếng Việt")

        path = self._pdf("unsupported.pdf", ["native", "scan"])
        with self.assertRaises(DocumentIngestionError) as caught:
            load_and_chunk_document(path, InvalidOCR(), self.settings, self.cache)
        self.assertEqual(caught.exception.failed_pages, [2])
        self.assertIn("thiếu ký tự tiếng Việt", str(caught.exception))

    def test_invalid_ocr_alphabet_preserves_short_native_text(self):
        class InvalidOCR(FakeOCR):
            def recognize_page(self, _image):
                raise OCRConfigurationError("Model thiếu ký tự tiếng Việt")

        path = self._pdf("short-supported.pdf", ["short"])
        result = load_and_chunk_document(path, InvalidOCR(), self.settings, self.cache)
        self.assertIn("Dieu 1", result.chunks[0].page_content)

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.cache = OCRResultCache(self.root / "cache")
        self.settings = IngestionSettings(
            ocr_enabled=True,
            dpi=96,
            native_text_threshold=50,
            sanity_min_chars=3,
            blank_pixel_threshold=245,
            blank_dark_ratio=0.0005,
            blank_stddev=2.0,
            chunk_size=1000,
            chunk_overlap=0,
            cache_path=self.root / "cache",
        )

    def tearDown(self):
        self.temporary.cleanup()

    def test_section_heading_is_inherited_by_following_chunks_on_page(self):
        chunks = [
            Document(
                page_content="2 3.3 Địa điểm III. NÓI DUNG - Giới thiệu tổng quan",
                metadata={"page": 1},
            ),
            Document(
                page_content="Hướng dẫn xem thời khóa biểu và đăng ký học phần",
                metadata={"page": 1},
            ),
            Document(
                page_content="IV. KINH PHÍ - Nhà trường hỗ trợ tổ chức",
                metadata={"page": 2},
            ),
        ]

        annotate_chunk_metadata(chunks)

        self.assertEqual(detect_section_title(chunks[0].page_content), "NỘI DUNG")
        self.assertEqual(chunks[0].metadata["section_title"], "NỘI DUNG")
        self.assertEqual(chunks[1].metadata["section_title"], "NỘI DUNG")
        self.assertEqual(chunks[1].metadata["page_chunk_index"], 1)
        self.assertEqual(chunks[2].metadata["section_title"], "KINH PHÍ")

    @staticmethod
    def _image_bytes():
        image = Image.new("RGB", (500, 180), "white")
        draw = ImageDraw.Draw(image)
        draw.rectangle((20, 30, 470, 140), outline="black", width=4)
        draw.text((35, 70), "DIEU 12 - NOI DUNG SCAN", fill="black")
        output = io.BytesIO()
        image.save(output, format="PNG")
        return output.getvalue()

    def _pdf(self, name, page_kinds):
        path = self.root / name
        with fitz.open() as document:
            for kind in page_kinds:
                page = document.new_page(width=500, height=220)
                if kind == "native":
                    page.insert_text(
                        (20, 40),
                        "Quy dinh hoc vu co noi dung native du dai de khong can goi OCR. "
                        "Thong tin nay chi dung cho unit test.",
                    )
                elif kind == "short":
                    page.insert_text((20, 40), "Dieu 1")
                elif kind == "scan":
                    page.insert_image(page.rect, stream=self._image_bytes())
                elif kind == "uncertain":
                    page.draw_rect(fitz.Rect(30, 30, 200, 120), color=(0, 0, 0), fill=(0, 0, 0))
                elif kind != "blank":
                    raise AssertionError(kind)
            document.save(path)
        return path

    def test_text_pdf_does_not_call_ocr(self):
        path = self._pdf("native.pdf", ["native"])
        fake = FakeOCR()
        result = load_and_chunk_document(path, fake, self.settings, self.cache)
        self.assertEqual(fake.calls, 0)
        self.assertEqual(result.chunks[0].metadata["extraction_method"], "native")

    def test_native_table_keeps_each_level_with_its_own_scores(self):
        path = self.root / "levels.pdf"
        with fitz.open() as document:
            page = document.new_page(width=580, height=240)
            x_positions = (20, 65, 155, 250, 350, 450, 560)
            y_positions = (25, 55, 85, 115)
            for x in x_positions:
                page.draw_line((x, 25), (x, 115), color=(0, 0, 0))
            for y in y_positions:
                page.draw_line((20, y), (560, y), color=(0, 0, 0))
            cells = (
                ("TT", "Language", "Certificate", "Bac 3", "Bac 4", "Bac 5"),
                ("1", "English", "TOEIC", "Listen 275-399", "Listen 400-489", "Listen 490-495"),
                ("", "", "IELTS", "4.0-5.0", "5.5-6.5", "7.0-8.0"),
            )
            for row_index, row in enumerate(cells):
                for column, value in enumerate(row):
                    if value:
                        page.insert_text((x_positions[column] + 4, y_positions[row_index] + 19), value, fontsize=8)
            page.insert_text((25, 150), "Other policy text outside the table must remain searchable.")
            document.save(path)

        result = load_and_chunk_document(path, settings=self.settings, cache=self.cache)
        toeic = [chunk for chunk in result.chunks if "TOEIC" in chunk.page_content]
        self.assertEqual(len(toeic), 3)
        bac3 = next(chunk for chunk in toeic if "Bac 3" in chunk.page_content)
        bac4 = next(chunk for chunk in toeic if "Bac 4" in chunk.page_content)
        bac5 = next(chunk for chunk in toeic if "Bac 5" in chunk.page_content)
        self.assertIn("Listen 275-399", bac3.page_content)
        self.assertNotIn("400-489", bac3.page_content)
        self.assertIn("Listen 400-489", bac4.page_content)
        self.assertNotIn("275-399", bac4.page_content)
        self.assertNotIn("490-495", bac4.page_content)
        self.assertIn("Listen 490-495", bac5.page_content)
        self.assertNotIn("400-489", bac5.page_content)
        self.assertEqual(bac4.metadata["extraction_method"], "native_table")
        self.assertEqual(bac4.metadata["page"], 0)
        self.assertEqual(bac4.metadata["table_level"], 4)
        self.assertTrue(any(
            chunk.metadata["extraction_method"] == "native"
            and "Other policy text" in chunk.page_content
            and "TOEIC" not in chunk.page_content
            for chunk in result.chunks
        ))

    def test_scan_pdf_calls_fake_ocr_and_keeps_zero_based_page(self):
        path = self._pdf("scan.pdf", ["scan"])
        fake = FakeOCR()
        result = load_and_chunk_document(path, fake, self.settings, self.cache)
        self.assertEqual(fake.calls, 1)
        self.assertEqual(result.chunks[0].metadata["page"], 0)
        self.assertEqual(result.chunks[0].metadata["extraction_method"], "ppocrv6")
        self.assertIsInstance(result.chunks[0].metadata["ocr_confidence"], float)

    def test_mixed_pdf_only_ocr_pages_without_enough_text(self):
        path = self._pdf("mixed.pdf", ["native", "scan", "native"])
        fake = FakeOCR()
        result = load_and_chunk_document(path, fake, self.settings, self.cache)
        self.assertEqual(fake.calls, 1)
        self.assertEqual(
            [chunk.metadata["extraction_method"] for chunk in result.chunks],
            ["native", "ppocrv6", "native"],
        )

    def test_short_native_text_survives_ocr_failure(self):
        path = self._pdf("short.pdf", ["short"])
        result = load_and_chunk_document(path, FakeOCR(fail=True), self.settings, self.cache)
        self.assertIn("Dieu 1", result.chunks[0].page_content)
        self.assertEqual(result.chunks[0].metadata["extraction_method"], "native_low_text")

    def test_blank_page_is_skipped(self):
        path = self._pdf("blank.pdf", ["blank", "native"])
        fake = FakeOCR()
        result = load_and_chunk_document(path, fake, self.settings, self.cache)
        self.assertEqual(fake.calls, 0)
        self.assertEqual(len(result.chunks), 1)
        self.assertTrue(any("trang trắng 1" in warning for warning in result.warnings))

    def test_uncertain_page_is_not_treated_as_blank(self):
        path = self._pdf("uncertain.pdf", ["uncertain"])
        fake = FakeOCR()
        load_and_chunk_document(path, fake, self.settings, self.cache)
        self.assertEqual(fake.calls, 1)

    def test_ocr_failure_lists_one_based_pages(self):
        path = self._pdf("failed.pdf", ["scan", "scan"])
        with self.assertRaises(DocumentIngestionError) as context:
            load_and_chunk_document(path, FakeOCR(fail=True), self.settings, self.cache)
        self.assertEqual(context.exception.failed_pages, [1, 2])
        self.assertIn("Trang thất bại: 1, 2", str(context.exception))

    def test_ocr_disabled_does_not_call_service(self):
        path = self._pdf("disabled.pdf", ["scan"])
        fake = FakeOCR()
        settings = IngestionSettings(**{**self.settings.__dict__, "ocr_enabled": False})
        with self.assertRaises(DocumentIngestionError):
            load_and_chunk_document(path, fake, settings, self.cache)
        self.assertEqual(fake.calls, 0)

    def test_docx_path_is_unchanged_and_never_calls_ocr(self):
        path = self.root / "document.docx"
        document = DocxDocument()
        document.add_paragraph(
            "Noi dung DOCX du dai de tiep tuc dung loader cu va khong can OCR. "
            "Day la du lieu unit test khong nhay cam."
        )
        document.save(path)
        fake = FakeOCR()
        result = load_and_chunk_document(path, fake, self.settings, self.cache)
        self.assertEqual(fake.calls, 0)
        self.assertEqual(result.chunks[0].metadata["extraction_method"], "native")

    def test_cache_hit_does_not_call_model_again(self):
        path = self._pdf("cache-hit.pdf", ["scan"])
        first = FakeOCR(fingerprint="same")
        load_and_chunk_document(path, first, self.settings, self.cache)
        second = FakeOCR(fingerprint="same", fail=True)
        load_and_chunk_document(path, second, self.settings, self.cache)
        self.assertEqual(first.calls, 1)
        self.assertEqual(second.calls, 0)

    @unittest.skipIf(os.name == "nt", "POSIX permission bits are not enforced on Windows")
    def test_cache_files_and_directories_are_private(self):
        path = self._pdf("private-cache.pdf", ["scan"])
        result = load_and_chunk_document(path, FakeOCR(), self.settings, self.cache)
        cache_file = result.created_cache_files[0]
        self.assertEqual(self.cache.root.stat().st_mode & 0o777, 0o700)
        self.assertEqual(cache_file.parent.stat().st_mode & 0o777, 0o700)
        self.assertEqual(cache_file.stat().st_mode & 0o777, 0o600)

    def test_fingerprint_change_creates_cache_miss(self):
        path = self._pdf("cache-miss.pdf", ["scan"])
        first = FakeOCR(fingerprint="config-a")
        second = FakeOCR(fingerprint="config-b")
        load_and_chunk_document(path, first, self.settings, self.cache)
        load_and_chunk_document(path, second, self.settings, self.cache)
        self.assertEqual(first.calls, 1)
        self.assertEqual(second.calls, 1)

    def test_delete_document_removes_only_its_cache(self):
        first_path = self._pdf("first.pdf", ["scan"])
        second_path = self._pdf("second.pdf", ["scan", "native"])
        load_and_chunk_document(first_path, FakeOCR(), self.settings, self.cache)
        load_and_chunk_document(second_path, FakeOCR(), self.settings, self.cache)
        from document_ingestion import sha256_file

        first_hash = sha256_file(first_path)
        second_hash = sha256_file(second_path)

        self.cache.delete_document(first_path)

        self.assertFalse((self.cache.root / first_hash).exists())
        self.assertTrue((self.cache.root / second_hash).exists())


if __name__ == "__main__":
    unittest.main()
