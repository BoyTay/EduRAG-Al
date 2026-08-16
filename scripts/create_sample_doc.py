import docx

doc = docx.Document()
doc.add_heading('QUY CHẾ ĐÀO TẠO & CÔNG TÁC SINH VIÊN (BẢN TÓM TẮT MẪU)', 0)

doc.add_heading('Điều 1: Đăng ký học phần và Số tín chỉ tối thiểu', level=1)
doc.add_paragraph(
    '1. Sinh viên hệ chính quy phải đăng ký tối thiểu 14 tín chỉ trong một học kỳ chính '
    '(trừ học kỳ cuối cùng chuẩn bị tốt nghiệp). Số tín chỉ tối đa sinh viên được phép đăng ký '
    'là 25 tín chỉ để đảm bảo chất lượng học tập.'
)
doc.add_paragraph(
    '2. Đối với học kỳ phụ (học kỳ hè), sinh viên được đăng ký tối đa 8 tín chỉ và không quy định '
    'số tín chỉ tối thiểu.'
)

doc.add_heading('Điều 2: Điều kiện xét học bổng khuyến khích học tập', level=1)
doc.add_paragraph(
    '1. Điều kiện xét học bổng khuyến khích học tập bao gồm:\n'
    '- Có điểm trung bình học kỳ (GPA) đạt từ 3.2 trở lên (theo thang điểm 4.0).\n'
    '- Điểm rèn luyện học kỳ đạt từ 80 điểm trở lên (loại Tốt trở lên).\n'
    '- Không bị kỷ luật từ mức khiển trách trở lên trong học kỳ xét học bổng.\n'
    '- Không có học phần nào bị điểm F (phải thi lại hoặc học lại) trong học kỳ đó.'
)
doc.add_paragraph(
    '2. Học bổng được cấp theo thứ tự từ cao xuống thấp cho đến khi hết quỹ học bổng của Khoa.'
)

doc.add_heading('Điều 3: Nghỉ học tạm thời và Bảo lưu kết quả học tập', level=1)
doc.add_paragraph(
    '1. Sinh viên được quyền xin nghỉ học tạm thời và bảo lưu kết quả học tập trong các trường hợp sau:\n'
    '- Được điều động vào lực lượng vũ trang (nghĩa vụ quân sự).\n'
    '- Bị ốm đau, tai nạn hoặc bệnh tật phải điều trị dài ngày (có xác nhận của cơ sở y tế cấp quận/huyện trở lên).\n'
    '- Vì nhu cầu cá nhân, tuy nhiên sinh viên phải học ít nhất 1 học kỳ tại trường và không thuộc diện bị buộc thôi học.'
)
doc.add_paragraph(
    '2. Thời gian nghỉ học tạm thời vì nhu cầu cá nhân không được quá 4 học kỳ chính.'
)

doc.add_heading('Điều 4: Điều kiện tốt nghiệp đại học', level=1)
doc.add_paragraph(
    '1. Sinh viên được xét công nhận tốt nghiệp khi tích lũy đủ số tín chỉ quy định của chương trình đào tạo.\n'
    '- Điểm trung bình tích lũy toàn khóa (CPA) đạt từ 2.00 trở lên (thang điểm 4.0).\n'
    '- Đạt chuẩn đầu ra về ngoại ngữ (ví dụ: TOEIC 450 hoặc tương đương) và tin học theo quy định của nhà trường.\n'
    '- Có chứng chỉ Giáo dục quốc phòng - An ninh và Giáo dục thể chất.\n'
    '- Không bị truy cứu trách nhiệm hình sự hoặc không trong thời gian bị kỷ luật ở mức đình chỉ học tập.'
)

doc.save('c:/Data/DoAnChuyenNganh/EduRAG Al/data/Quy_che_dao_tao_mau.docx')
print('Successfully created sample document!')
